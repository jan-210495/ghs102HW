import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

d = docx.Document()
p = d.add_paragraph("نص عربي ")
r = p.add_run()
rPr = r._r.get_or_add_rPr()
st = OxmlElement('w:rStyle'); st.set(qn('w:val'), 'FootnoteReference'); rPr.append(st)
fn = OxmlElement('w:footnoteReference'); fn.set(qn('w:id'), '1'); r._r.append(fn)
d.save('/tmp/t1.docx')
