# -*- coding: utf-8 -*-
"""
محرك بناء ملف وورد عربي (RTL) لوظيفة «الرأي العام».
يدعم: الهوامش السفلية الحقيقية، حقل فهرس آلي (TOC)، حقول PAGEREF،
جداول RTL، أشكالاً، ترويسة وتذييلاً مع ترقيم صفحات.
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

FONT = "Traditional Arabic"
FONT_SIZE = 14          # المتن
FN_SIZE = 11            # الهوامش

PPR_AFTER_BIDI = ('w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
                  'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap',
                  'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
                  'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
                  'w:pPrChange')

RPR_AFTER_RTL = ('w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish', 'w:oMath')

RPR_SEQ = ('w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps', 'w:smallCaps',
           'w:strike', 'w:dstrike', 'w:outline', 'w:shadow', 'w:emboss', 'w:imprint',
           'w:noProof', 'w:snapToGrid', 'w:vanish', 'w:webHidden', 'w:color', 'w:spacing',
           'w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect',
           'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang',
           'w:eastAsianLayout', 'w:specVanish', 'w:oMath')

TRPR_SEQ = ('w:cnfStyle', 'w:divId', 'w:gridBefore', 'w:gridAfter', 'w:wBefore', 'w:wAfter',
            'w:cantSplit', 'w:trHeight', 'w:tblHeader', 'w:tblCellSpacing', 'w:jc', 'w:hidden')


TBLPR_AFTER_BIDI = ('w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
                    'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd', 'w:tblLayout',
                    'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')


def table_rtl(t):
    """ضبط الجدول من اليمين إلى اليسار مع احترام ترتيب عناصر tblPr."""
    tblPr = t._tbl.tblPr
    if tblPr.find(qn('w:bidiVisual')) is None:
        tblPr.insert_element_before(OxmlElement('w:bidiVisual'), *TBLPR_AFTER_BIDI)
    return t


def rpr_set(rPr, tag, val=None):
    """إدراج عنصر داخل rPr بالترتيب الذي يفرضه مخطط OOXML."""
    el = rPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        idx = RPR_SEQ.index(tag)
        rPr.insert_element_before(el, *RPR_SEQ[idx + 1:])
    if val is not None:
        el.set(qn('w:val'), str(val))
    return el


def trpr_set(trPr, tag, val=None):
    el = trPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        idx = TRPR_SEQ.index(tag)
        trPr.insert_element_before(el, *TRPR_SEQ[idx + 1:])
    if val is not None:
        el.set(qn('w:val'), str(val))
    return el


# ---------------------------------------------------------------- أدوات XML
def _insert_before(parent, elm, successors):
    parent.insert_element_before(elm, *successors)
    return elm


def para_rtl(p, align='both', indent=None, line=1.5, before=0, after=0,
             keep=False, bidi=True):
    pPr = p._p.get_or_add_pPr()
    if bidi:
        b = OxmlElement('w:bidi')
        if pPr.find(qn('w:bidi')) is None:
            _insert_before(pPr, b, PPR_AFTER_BIDI)
    pf = p.paragraph_format
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if line is not None:
        pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if keep:
        pf.keep_with_next = True
    pf.alignment = {'both': WD_ALIGN_PARAGRAPH.JUSTIFY,
                    'right': WD_ALIGN_PARAGRAPH.RIGHT,
                    'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    return p


def run_rtl(r, size=None, bold=None, italic=None, color=None, font=None):
    rPr = r._r.get_or_add_rPr()
    f = font or FONT
    r.font.name = f
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:cs'), f)
    rFonts.set(qn('w:ascii'), f)
    rFonts.set(qn('w:hAnsi'), f)
    if size:
        r.font.size = Pt(size)
        rpr_set(rPr, 'w:szCs', int(size * 2))
    if bold is not None:
        rpr_set(rPr, 'w:b', '1' if bold else '0')
        rpr_set(rPr, 'w:bCs', '1' if bold else '0')
    if italic is not None:
        rpr_set(rPr, 'w:i', '1' if italic else '0')
        rpr_set(rPr, 'w:iCs', '1' if italic else '0')
    if color:
        r.font.color.rgb = color
    rpr_set(rPr, 'w:rtl', '1')
    return r


DIACRITICS = '\u064B\u064C\u064D\u064E\u064F\u0650\u0651\u0652\u0653\u0654\u0655\u0670\u06D6\u06DC\u06DF\u06E0\u06E5\u06E6\u06E7\u06E8\u06EA\u06EB\u06EC\u06ED'
_DIAC_RE = re.compile('[' + DIACRITICS + ']')


_TANWEEN = '\u064B'
_PROTECT = '\u0001'


def clean(t):
    """إزالة التشكيل من النص المطبوع، مع الحفاظ على تنوين الفتح بعد الألف (الكتابة المعيارية)."""
    if not isinstance(t, str):
        return t
    # حماية تنوين الفتح الواقع بعد ألف، لأنه جزء من الرسم الإملائي لا زينة
    t = t.replace('\u0627' + _TANWEEN, '\u0627' + _PROTECT)
    t = _DIAC_RE.sub('', t)
    t = t.replace(_PROTECT, _TANWEEN)
    t = t.replace('\u200f', '').replace('\u200e', '')
    t = re.sub(r'[ ]{2,}', ' ', t)
    t = t.replace(' ,', '،').replace(' ،', '،')
    t = t.replace('؟', '؟').replace(' ?', '؟')
    return t


# التسطير السفلي لا يُستعمل علامةً للمائل إلا إذا بدأ النص بعده بحرف عربي،
# حمايةً للأرقام الجامعية مثل jan_436439 من التفسير الخاطئ.
TOKEN_RE = re.compile(r'\{\{(.+?)\}\}|\*\*(.+?)\*\*|_([\u0600-\u06FF][^_]*?)_', re.S)
INNER_RE = re.compile(r'\*\*(.+?)\*\*|_([\u0600-\u06FF][^_]*?)_', re.S)


class ArabicDocx:
    def __init__(self):
        self.doc = Document()
        self.footnotes = []          # قائمة نصوص الهوامش
        self._bm_id = 100
        self._setup_page()
        self._setup_styles()

    # ------------------------------------------------------------ إعداد الصفحة
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        s.header_distance = Cm(1.25)
        s.footer_distance = Cm(1.25)
        sectPr = s._sectPr
        # اتجاه المقطع من اليمين إلى اليسار
        bidi = OxmlElement('w:bidi')
        sectPr.insert_element_before(bidi, 'w:rtlGutter', 'w:docGrid', 'w:printerSettings')
        # موضع الهوامش السفلية أسفل الصفحة
        fpr = OxmlElement('w:footnotePr')
        pos = OxmlElement('w:pos')
        pos.set(qn('w:val'), 'pageBottom')
        fpr.append(pos)
        sectPr.insert_element_before(fpr, 'w:endnotePr', 'w:type', 'w:pgSz', 'w:pgMar',
                                     'w:paperSrc', 'w:pgBorders', 'w:lnNumType',
                                     'w:pgNumType', 'w:cols', 'w:formProt', 'w:vAlign',
                                     'w:noEndnote', 'w:titlePg', 'w:textDirection',
                                     'w:bidi', 'w:rtlGutter', 'w:docGrid')
        self.section = s

    def _style_rtl(self, style, size, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   before=0, after=0, line=1.5, indent=None, keep=False):
        pPr = style.element.get_or_add_pPr()
        if pPr.find(qn('w:bidi')) is None:
            _insert_before(pPr, OxmlElement('w:bidi'), PPR_AFTER_BIDI)
        f = style.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor(0, 0, 0)
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:cs'), FONT)
        rFonts.set(qn('w:ascii'), FONT)
        rFonts.set(qn('w:hAnsi'), FONT)
        rpr_set(rPr, 'w:szCs', int(size * 2))
        rpr_set(rPr, 'w:bCs', 1 if bold else 0)
        pf = style.paragraph_format
        pf.alignment = align
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        if indent is not None:
            pf.first_line_indent = Cm(indent)
        if keep:
            pf.keep_with_next = True

    def _setup_styles(self):
        st = self.doc.styles
        self._style_rtl(st['Normal'], FONT_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.0)
        self._style_rtl(st['Heading 1'], 18, bold=True,
                        align=WD_ALIGN_PARAGRAPH.RIGHT, before=12, after=8, line=1.3, keep=True)
        self._style_rtl(st['Heading 2'], 16, bold=True,
                        align=WD_ALIGN_PARAGRAPH.RIGHT, before=10, after=6, line=1.3, keep=True)
        if 'Heading 3' in [x.name for x in st]:
            self._style_rtl(st['Heading 3'], 14, bold=True,
                            align=WD_ALIGN_PARAGRAPH.RIGHT, before=8, after=4, line=1.3, keep=True)
        # نمط الهوامش
        styles_el = st.element
        for sid, name, tag, sz, vaa in (
                ('FootnoteText', 'footnote text', 'w:pStyle', FN_SIZE, None),
                ('FootnoteReference', 'footnote reference', None, None, 'superscript')):
            if styles_el.find(f'.//{{*}}style[@{{*}}styleId="{sid}"]') is not None:
                continue
            s = OxmlElement('w:style')
            s.set(qn('w:type'), 'character' if sid == 'FootnoteReference' else 'paragraph')
            s.set(qn('w:styleId'), sid)
            nm = OxmlElement('w:name')
            nm.set(qn('w:val'), name)
            s.append(nm)
            if sid == 'FootnoteText':
                base = OxmlElement('w:basedOn')
                base.set(qn('w:val'), 'Normal')
                s.append(base)
                pPr = OxmlElement('w:pPr')
                bidi = OxmlElement('w:bidi')
                pPr.append(bidi)
                jc = OxmlElement('w:jc')
                jc.set(qn('w:val'), 'both')
                pPr.append(jc)
                sp = OxmlElement('w:spacing')
                sp.set(qn('w:before'), '0')
                sp.set(qn('w:after'), '0')
                sp.set(qn('w:line'), '240')
                sp.set(qn('w:lineRule'), 'auto')
                pPr.append(sp)
                s.append(pPr)
                rPr = OxmlElement('w:rPr')
                for t, v in (('w:sz', sz * 2), ('w:szCs', sz * 2)):
                    e = OxmlElement(t)
                    e.set(qn('w:val'), str(v))
                    rPr.append(e)
                s.append(rPr)
            else:
                rPr = OxmlElement('w:rPr')
                va = OxmlElement('w:vertAlign')
                va.set(qn('w:val'), vaa)
                rPr.append(va)
                s.append(rPr)
            styles_el.append(s)

    def new_section(self, restart_numbering=True):
        from docx.enum.section import WD_SECTION
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        s.header_distance = Cm(1.25)
        s.footer_distance = Cm(1.25)
        sectPr = s._sectPr
        if sectPr.find(qn('w:bidi')) is None:
            sectPr.insert_element_before(OxmlElement('w:bidi'), 'w:rtlGutter',
                                         'w:docGrid', 'w:printerSettings')
        fpr = OxmlElement('w:footnotePr')
        pos = OxmlElement('w:pos')
        pos.set(qn('w:val'), 'pageBottom')
        fpr.append(pos)
        sectPr.insert_element_before(fpr, 'w:endnotePr', 'w:type', 'w:pgSz', 'w:pgMar',
                                     'w:paperSrc', 'w:pgBorders', 'w:lnNumType',
                                     'w:pgNumType', 'w:cols', 'w:formProt', 'w:vAlign',
                                     'w:noEndnote', 'w:titlePg', 'w:textDirection',
                                     'w:bidi', 'w:rtlGutter', 'w:docGrid')
        if restart_numbering:
            pg = OxmlElement('w:pgNumType')
            pg.set(qn('w:start'), '1')
            sectPr.insert_element_before(pg, 'w:cols', 'w:formProt', 'w:vAlign',
                                         'w:noEndnote', 'w:titlePg', 'w:textDirection',
                                         'w:bidi', 'w:rtlGutter', 'w:docGrid')
        self.section = s
        return s

    # ------------------------------------------------------------ ترويسة وتذييل
    def add_header_footer(self, header_text):
        sec = self.section
        sec.different_first_page_header_footer = True
        h = sec.header
        p = h.paragraphs[0]
        p.text = ''
        para_rtl(p, align='right', line=1.0, after=0)
        run_rtl(p.add_run(header_text), size=10, color=RGBColor(0x44, 0x44, 0x44))
        # خط أسفل الترويسة
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '999999')
        pBdr.append(bottom)
        pPr.insert_element_before(pBdr, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens',
                                  'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                                  'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
                                  'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
                                  'w:spacing', 'w:ind', 'w:contextualSpacing',
                                  'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
                                  'w:textDirection', 'w:textAlignment', 'w:rPr')
        f = sec.footer
        fp = f.paragraphs[0]
        fp.text = ''
        para_rtl(fp, align='center', line=1.0)
        r = fp.add_run()
        run_rtl(r, size=11)
        self._add_field(r, 'PAGE', '1', dirty=True)
        self._add_field_text_around(fp, '')

    def _add_field_text_around(self, p, txt):
        if txt:
            run_rtl(p.add_run(txt), size=11)

    # ------------------------------------------------------------ الحقول
    def _add_field(self, run, instr, placeholder='', dirty=True):
        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        if dirty:
            fld_begin.set(qn('w:dirty'), 'true')
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = f' {instr} '
        sep = OxmlElement('w:fldChar')
        sep.set(qn('w:fldCharType'), 'separate')
        t = OxmlElement('w:t')
        t.text = placeholder
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        for el in (fld_begin, it, sep, t, end):
            run._r.append(el)

    def toc_field(self, levels='1-3'):
        p = self.doc.add_paragraph()
        para_rtl(p, align='both', line=1.3)
        r = p.add_run()
        run_rtl(r, size=FONT_SIZE)
        self._add_field(r, f'TOC \\o "{levels}" \\h \\z \\u',
                        'سيُبنى هذا الفهرس تلقائياً عند فتح الملف في برنامج وورد، أو بتحديث الحقول من قائمة المراجع.',
                        dirty=True)
        self._enable_update_fields()

    def pageref_field(self, p, bookmark, size=FONT_SIZE):
        r = p.add_run()
        run_rtl(r, size=size)
        self._add_field(r, f'PAGEREF {bookmark} \\h', '—', dirty=True)

    def _enable_update_fields(self):
        s = self.doc.settings.element
        if s.find(qn('w:updateFields')) is None:
            uf = OxmlElement('w:updateFields')
            uf.set(qn('w:val'), 'true')
            s.insert_element_before(uf, 'w:hdrShapeDefaults', 'w:footnotePr', 'w:endnotePr',
                                    'w:compat', 'w:docVars', 'w:rsids', 'w:attachedSchema',
                                    'w:themeFontLang', 'w:clrSchemeMapping',
                                    'w:doNotIncludeSubdocsInStats', 'w:doNotAutoCompressPictures',
                                    'w:forceUpgrade', 'w:captions', 'w:readModeInkLockDown',
                                    'w:smartTagType', 'w:schemaLibrary', 'w:shapeDefaults',
                                    'w:doNotEmbedSmartTags', 'w:decimalSymbol', 'w:listSeparator',
                                    'w:docId', 'w:defaultImageDpi')

    def bookmark(self, p, name):
        self._bm_id += 1
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(self._bm_id))
        start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(self._bm_id))
        pPr = p._p.find(qn('w:pPr'))
        if pPr is not None:
            pPr.addnext(start)
        else:
            p._p.insert(0, start)
        p._p.append(end)

    # ------------------------------------------------------------ فقرات
    def _add_text(self, p, text, bold=False, italic=False, size=None, color=None):
        if not text:
            return
        r = p.add_run(clean(text))
        run_rtl(r, size=size or FONT_SIZE, bold=bold, italic=italic, color=color)

    def _add_footnote(self, p, text):
        self.footnotes.append(clean(text))
        fid = len(self.footnotes)
        r = p.add_run()
        rPr = r._r.get_or_add_rPr()
        stl = OxmlElement('w:rStyle')
        stl.set(qn('w:val'), 'FootnoteReference')
        rPr.append(stl)
        run_rtl(r, size=FONT_SIZE)
        ref = OxmlElement('w:footnoteReference')
        ref.set(qn('w:id'), str(fid))
        r._r.append(ref)

    def para(self, text, align='both', indent=1.0, line=1.5, before=0, after=6,
             bold=False, italic=False, size=None, keep=False, style=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        para_rtl(p, align=align, indent=indent, line=line, before=before, after=after, keep=keep)
        pos = 0
        for m in TOKEN_RE.finditer(text):
            if m.start() > pos:
                self._add_text(p, text[pos:m.start()], bold=bold, italic=italic, size=size)
            if m.group(1):
                self._add_footnote(p, m.group(1))
            elif m.group(2):
                self._add_text(p, m.group(2), bold=True, italic=italic, size=size)
            else:
                self._add_text(p, m.group(3), bold=bold, italic=True, size=size)
            pos = m.end()
        if pos < len(text):
            self._add_text(p, text[pos:], bold=bold, italic=italic, size=size)
        return p

    def heading(self, text, level=1, page_break=False):
        p = self.doc.add_paragraph(style=f'Heading {level}')
        para_rtl(p, align='right', indent=0, line=1.3,
                 before={1: 0 if page_break else 12, 2: 10, 3: 8}[level],
                 after={1: 8, 2: 6, 3: 4}[level], keep=True)
        if page_break:
            p.paragraph_format.page_break_before = True
        self._add_text(p, text, bold=True,
                       size={1: 18, 2: 16, 3: 14}[level])
        return p

    def caption(self, text, bookmark=None, above=True):
        p = self.doc.add_paragraph()
        para_rtl(p, align='center', indent=0, line=1.15, before=6 if above else 4,
                 after=4 if above else 10, keep=True)
        self._add_text(p, text, bold=True, size=12)
        if bookmark:
            self.bookmark(p, bookmark)
        return p

    def page_break(self):
        self.doc.add_page_break()

    # ------------------------------------------------------------ جداول
    def table(self, data, widths=None, header_rows=1, font_size=11,
              align_center_cols=None, caption=None, bookmark=None):
        if caption:
            self.caption(caption, bookmark=bookmark, above=True)
        rows = len(data)
        cols = len(data[0])
        t = self.doc.add_table(rows=rows, cols=cols)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_rtl(t)
        # تكرار سطر العنوان
        for ri, row in enumerate(data):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            if ri < header_rows:
                trpr_set(trPr, 'w:tblHeader')
            trpr_set(trPr, 'w:cantSplit')
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ''
                p = cell.paragraphs[0]
                para_rtl(p, align='center' if (ri < header_rows or (align_center_cols and ci in align_center_cols)) else 'both',
                         indent=0, line=1.15, before=2, after=2)
                pos = 0
                for m in TOKEN_RE.finditer(str(val)):
                    if m.start() > pos:
                        self._add_text(p, str(val)[pos:m.start()],
                                       bold=(ri < header_rows), size=font_size)
                    if m.group(1):
                        self._add_footnote(p, m.group(1))
                    elif m.group(2):
                        self._add_text(p, m.group(2), bold=True, size=font_size)
                    else:
                        self._add_text(p, m.group(3), italic=True, size=font_size)
                    pos = m.end()
                if pos < len(str(val)):
                    self._add_text(p, str(val)[pos:], bold=(ri < header_rows), size=font_size)
                if ri < header_rows:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:fill'), 'E8E8E8')
                    tcPr.append(shd)
        if widths:
            for ci, w in enumerate(widths):
                for row in t.rows:
                    row.cells[ci].width = Cm(w)
        # مسافة بعد الجدول
        p = self.doc.add_paragraph()
        para_rtl(p, align='both', line=1.0, after=0)
        return t

    def figure(self, path, width_cm=15.0, caption=None, bookmark=None):
        p = self.doc.add_paragraph()
        para_rtl(p, align='center', indent=0, line=1.0, before=6, after=2, keep=True)
        r = p.add_run()
        r.add_picture(path, width=Cm(width_cm))
        if caption:
            self.caption(caption, bookmark=bookmark, above=False)

    def pagebreak_para(self):
        p = self.doc.add_paragraph()
        para_rtl(p, align='both', line=1.0)
        r = p.add_run()
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        r._r.append(br)

    # ------------------------------------------------------------ ضبط المقاطع
    def finalize_sections(self):
        """ضبط مراجع الترويسة والتذييل وترقيم الصفحات في المقاطع نهائياً."""
        body = self.doc.element.body
        sentinel = body.find(qn('w:sectPr'))
        para_sects = [p.find(qn('w:pPr')).find(qn('w:sectPr'))
                      for p in body.iter(qn('w:p'))
                      if p.find(qn('w:pPr')) is not None
                      and p.find(qn('w:pPr')).find(qn('w:sectPr')) is not None]
        all_sects = para_sects + ([sentinel] if sentinel is not None else [])

        # 1) إزالة تكرار footnotePr داخل المقطع نفسه
        for sp in all_sects:
            seen = False
            for el in list(sp):
                if el.tag == qn('w:footnotePr'):
                    if seen:
                        sp.remove(el)
                    seen = True

        # 2) نقل مراجع الترويسة والتذييل إلى مقطع الجسم فقط
        refs = []
        for sp in all_sects:
            for tag in ('w:headerReference', 'w:footerReference'):
                for el in list(sp.findall(qn(tag))):
                    if sp is not sentinel:
                        sp.remove(el)
                        refs.append(el)
        if sentinel is not None and refs:
            for el in refs:
                if el.tag == qn('w:headerReference'):
                    if sentinel.find(qn('w:headerReference')) is None:
                        sentinel.insert(0, el)
                else:
                    if sentinel.find(qn('w:footerReference')) is None:
                        idx = 1 if sentinel.find(qn('w:headerReference')) is not None else 0
                        sentinel.insert(idx, el)

        # 3) إظهار الترويسة والتذييل في كل صفحات الجسم
        if sentinel is not None:
            for el in list(sentinel.findall(qn('w:titlePg'))):
                sentinel.remove(el)
            for el in list(sentinel.findall(qn('w:evenAndOddHeaders'))):
                sentinel.remove(el)

        # 4) مقطع الغلاف بلا ترقيم ولا ترويسة
        for sp in para_sects:
            for tag in ('w:pgNumType', 'w:titlePg'):
                for el in list(sp.findall(qn(tag))):
                    sp.remove(el)

    # ------------------------------------------------------------ الحفظ
    def _footnotes_xml(self):
        w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        head = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                f'<w:footnotes xmlns:w="{w}">')
        sep = ('<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" '
               'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr>'
               '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>')
        out = [head, sep]
        for i, txt in enumerate(self.footnotes, start=1):
            runs = []
            for m in INNER_RE.finditer(txt):
                pre = txt[:m.start()] if m.start() == 0 else ''
                if m.group(1):
                    runs.append(('b', m.group(1)))
                else:
                    runs.append(('i', m.group(2)))
            # بناء مبسط: نقسم النص إلى أجزاء
            parts, pos = [], 0
            for m in INNER_RE.finditer(txt):
                if m.start() > pos:
                    parts.append(('n', txt[pos:m.start()]))
                if m.group(1):
                    parts.append(('b', m.group(1)))
                else:
                    parts.append(('i', m.group(2)))
                pos = m.end()
            if pos < len(txt):
                parts.append(('n', txt[pos:]))
            r_xml = []
            for kind, seg in parts:
                seg = (seg.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))
                props = ''
                if kind == 'i':
                    props = ('<w:rPr><w:i/><w:iCs/><w:rFonts w:ascii="Traditional Arabic" '
                             'w:hAnsi="Traditional Arabic" w:cs="Traditional Arabic"/>'
                             '<w:sz w:val="22"/><w:szCs w:val="22"/><w:rtl/></w:rPr>')
                elif kind == 'b':
                    props = ('<w:rPr><w:b/><w:bCs/><w:rFonts w:ascii="Traditional Arabic" '
                             'w:hAnsi="Traditional Arabic" w:cs="Traditional Arabic"/>'
                             '<w:sz w:val="22"/><w:szCs w:val="22"/><w:rtl/></w:rPr>')
                else:
                    props = ('<w:rPr><w:rFonts w:ascii="Traditional Arabic" '
                             'w:hAnsi="Traditional Arabic" w:cs="Traditional Arabic"/>'
                             '<w:sz w:val="22"/><w:szCs w:val="22"/><w:rtl/></w:rPr>')
                r_xml.append(f'<w:r>{props}<w:t xml:space="preserve">{seg}</w:t></w:r>')
            out.append(
                f'<w:footnote w:id="{i}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/>'
                '<w:keepLines/><w:bidi/><w:jc w:val="both"/></w:pPr>'
                '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
                '<w:r><w:rPr><w:rFonts w:ascii="Traditional Arabic" w:hAnsi="Traditional Arabic" '
                'w:cs="Traditional Arabic"/><w:sz w:val="22"/><w:szCs w:val="22"/><w:rtl/></w:rPr>'
                '<w:t xml:space="preserve"> </w:t></w:r>'
                + ''.join(r_xml) + '</w:p></w:footnote>')
        out.append('</w:footnotes>')
        return ''.join(out).encode('utf-8')

    def save(self, path):
        self.finalize_sections()
        blob = self._footnotes_xml()
        part = Part(PackURI('/word/footnotes.xml'),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
                    blob, self.doc.part.package)
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        self._enable_update_fields()
        self.doc.save(path)
