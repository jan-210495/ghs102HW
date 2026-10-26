# -*- coding: utf-8 -*-
"""بناء الوظيفة الأكاديمية بصيغ: DOCX (الرئيسية)، Markdown، وHTML قابل للطباعة.

الاستعمال:  python3 tools/build.py
"""

import io
import base64
import os
import re
import sys
import html as html_lib
from collections import OrderedDict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from content_1 import BLOCKS as PART1
from content_2 import BLOCKS as PART2
from references import REFERENCES

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLOCKS = PART1 + PART2
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "الوظيفة_الأولى_الرأي_العام")
AR_FONT = "Traditional Arabic"

# =====================================================================
# أدوات مشتركة
# =====================================================================

CITE_RE = re.compile(r"\[\[(.*?)\]\]", re.S)


def resolve(text):
    """تحويل وسوم التوثيق إلى الصيغة (المؤلف، السنة، الصفحة)."""
    def repl(m):
        inner = m.group(1)
        outs = []
        for part in [x.strip() for x in inner.split(";;") if x.strip()]:
            if "|" in part:
                key, page = [x.strip() for x in part.split("|", 1)]
            else:
                key, page = part, ""
            if key not in REFERENCES:
                raise KeyError("مرجع غير معروف في المتن: " + key)
            year = REFERENCES[key]["year"]
            outs.append(key + "، " + year + ("، " + page if page else ""))
        return "(" + "؛ ".join(outs) + ")"
    return CITE_RE.sub(repl, text)


def plain(text):
    """نص بلا وسوم توثيق (لحساب الكلمات)."""
    return CITE_RE.sub("", text)


def ar_words(text):
    return re.findall(r"[\u0600-\u06ff\u0750-\u077f]+", plain(text))


def block_text(b):
    parts = []
    for k in ("text", "source", "label", "caption"):
        if isinstance(b.get(k), str):
            parts.append(b[k])
    for k in ("items", "lines"):
        v = b.get(k)
        if isinstance(v, list):
            parts.extend(str(x) for x in v)
    for k in ("headers", "rows"):
        v = b.get(k)
        if isinstance(v, list):
            for x in v:
                if isinstance(x, str):
                    parts.append(x)
                elif isinstance(x, (list, tuple)):
                    parts.extend(str(y) for y in x)
    return " ".join(parts)


def validate():
    used = set()
    for b in BLOCKS:
        for m in CITE_RE.finditer(block_text(b)):
            for part in [x.strip() for x in m.group(1).split(";;") if x.strip()]:
                key = part.split("|", 1)[0].strip()
                used.add(key)
                if key not in REFERENCES:
                    raise KeyError("مفتاح مرجع غير موجود: " + key)
    missing = set(REFERENCES) - used
    return used, missing


def segment_indices():
    jan_start = next(i for i, b in enumerate(BLOCKS)
                     if b["t"] == "h1" and b["text"].startswith("2. المقدمة"))
    wadee_start = next(i for i, b in enumerate(BLOCKS)
                       if b["t"] == "h1" and b["text"] == "القسم الثاني")
    refs_start = next(i for i, b in enumerate(BLOCKS)
                      if b["t"] == "h1" and b["text"] == "قائمة المصادر والمراجع")
    return jan_start, wadee_start, refs_start


def word_counts():
    js, ws, rs = segment_indices()
    jan = sum(len(ar_words(block_text(b))) for b in BLOCKS[js:ws])
    wadee = sum(len(ar_words(block_text(b))) for b in BLOCKS[ws:rs])
    return jan, wadee


def prose_counts():
    """كلمات النص المتصل فقط (الفقرات السردية دون الجداول والقوائم)."""
    js, ws, rs = segment_indices()
    jan = sum(len(ar_words(b["text"])) for b in BLOCKS[js:ws] if b["t"] == "p")
    wadee = sum(len(ar_words(b["text"])) for b in BLOCKS[ws:rs] if b["t"] == "p")
    return jan, wadee


def norm_ar(s):
    s = re.sub(r"^ال", "", s)
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = s.replace("ة", "ه").replace("ى", "ي").replace("ؤ", "و")
    s = re.sub(r"[\u064b-\u0652\u0640]", "", s)
    return s.strip()


def sorted_refs():
    ar = [(k, v) for k, v in REFERENCES.items() if v["lang"] == "ar"]
    en = [(k, v) for k, v in REFERENCES.items() if v["lang"] == "en"]
    ar.sort(key=lambda kv: norm_ar(kv[0]))
    en.sort(key=lambda kv: kv[1]["entry"].lower())
    return ar, en


def ar_num(n):
    return str(n).translate(str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩"))


# =====================================================================
# بناء ملف Word
# =====================================================================

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd, TC_AFTER_SHD)


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        insert_ordered(tcPr, borders, TC_AFTER_BORDERS)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for k, v in attrs.items():
            el.set(qn("w:" + k), str(v))


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        insert_ordered(pPr, OxmlElement("w:bidi"), PPR_AFTER_BIDI)
    return pPr


def style_run(run, size=13, bold=False, color=None, italic=False, font=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    name = font or AR_FONT
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        insert_ordered(rPr, lang, RPR_AFTER_LANG)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")
    return run


PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                  "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
                  "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
                  "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                  "w:autoSpaceDN", "w:bidi") + PPR_AFTER_BIDI
TC_AFTER_BORDERS = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                    "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
                    "w:cellMerge", "w:tcPrChange")
TC_AFTER_SHD = TC_AFTER_BORDERS[1:]
SECT_AFTER_BIDI = ("w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
SETTINGS_AFTER_UPDATE = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                         "w:compat", "w:docVars", "w:rsids", "m:mathPr",
                         "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                         "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures",
                         "w:forceUpgrade", "w:captions", "w:readModeInkLockDown",
                         "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults",
                         "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")
RPR_AFTER_LANG = ("w:eastAsianLayout", "w:specVanish", "w:oMath")


def _tagset(names):
    out = set()
    for n in names:
        try:
            out.add(qn(n))
        except KeyError:
            pass
    return out


def insert_ordered(parent, el, successors):
    """إدراج عنصر XML في موضعه الصحيح بحسب ترتيب مخطط OOXML."""
    succ = _tagset(successors)
    for child in parent:
        if child.tag in succ:
            child.addprevious(el)
            return
    parent.append(el)


def add_paragraph(doc, text, size=12.5, bold=False, color=None, align="right",
                  indent=True, space_before=0, space_after=5, line=1.2,
                  keep_next=False):
    para = doc.add_paragraph()
    set_rtl(para, {"right": WD_ALIGN_PARAGRAPH.RIGHT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align])
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Cm(0.65)
    if keep_next:
        pf.keep_with_next = True
    style_run(para.add_run(resolve(text)), size=size, bold=bold, color=color)
    return para


def add_field(paragraph, instr, size=9.5):
    """إضافة حقل Word مع نتيجة مخزّنة مؤقتاً كي يظهر رقم الصفحة في جميع العارضات."""
    r1 = paragraph.add_run()
    style_run(r1, size=size, color="555555")
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    r1._r.append(f1)

    r2 = paragraph.add_run()
    style_run(r2, size=size, color="555555")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " " + instr + " "
    r2._r.append(it)

    r3 = paragraph.add_run()
    style_run(r3, size=size, color="555555")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)

    r4 = paragraph.add_run("1")
    style_run(r4, size=size, color="555555")

    r5 = paragraph.add_run()
    style_run(r5, size=size, color="555555")
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    r5._r.append(f2)


def add_hr(doc, color="1F4E79"):
    para = doc.add_paragraph()
    set_rtl(para, WD_ALIGN_PARAGRAPH.CENTER)
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(8)
    pPr = para._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), color)
    bottom.set(qn("w:space"), "1")
    pbdr.append(bottom)
    insert_ordered(pPr, pbdr, PPR_AFTER_PBDR)
    style_run(para.add_run(" "), size=2)
    return para


def add_table(doc, headers, rows, widths=None, caption=None, fontsize=10):
    if caption:
        add_paragraph(doc, caption, size=10.5, bold=True, color="1F4E79",
                      align="center", indent=False, space_before=6,
                      space_after=4, keep_next=True)
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    tblPr.get_or_add_bidiVisual()
    total = 16.4  # عرض النص التقريبي بالسنتيمتر
    if widths:
        cw = [Cm(total * w / sum(widths)) for w in widths]
    else:
        cw = [Cm(total / len(headers))] * len(headers)
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = cw[i]
        set_cell_shading(hdr[i], "DCE6F1")
        p = hdr[i].paragraphs[0]
        set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        style_run(p.add_run(resolve(str(h))), size=fontsize, bold=True, color="1F4E79")
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = cw[i]
            if ri % 2 == 1:
                set_cell_shading(cells[i], "F5F8FC")
            p = cells[i].paragraphs[0]
            set_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            style_run(p.add_run(resolve(str(val))), size=fontsize)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def add_box(doc, text, label=None, kind="note"):
    fill = "EAF1F8" if kind == "note" else "F7F4EA"
    edge = "2E74B5" if kind == "note" else "BF9000"
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, right={"val": "single", "sz": "24", "color": edge})
    if label:
        p = cell.paragraphs[0]
        set_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        p.paragraph_format.space_after = Pt(3)
        style_run(p.add_run(label + " : "), size=11.5, bold=True, color=edge)
        p2 = cell.add_paragraph()
    else:
        p2 = cell.paragraphs[0]
    set_rtl(p2, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p2.paragraph_format.line_spacing = 1.2
    p2.paragraph_format.space_after = Pt(2)
    style_run(p2.add_run(resolve(text)), size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.1)
    sectPr = sec._sectPr
    if sectPr.find(qn("w:bidi")) is None:
        insert_ordered(sectPr, OxmlElement("w:bidi"), SECT_AFTER_BIDI)

    # إعدادات الخط الافتراضية واللغة
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = AR_FONT
    normal.font.size = Pt(12.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rfonts.set(qn("w:cs"), AR_FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")
    insert_ordered(rpr, lang, RPR_AFTER_LANG)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)

    # الترويسة والتذييل
    hp = sec.header.paragraphs[0]
    set_rtl(hp, WD_ALIGN_PARAGRAPH.CENTER)
    hp.paragraph_format.space_after = Pt(2)
    style_run(hp.add_run("الجامعة الافتراضية السورية (SVU)  |  المدخل إلى علم الاجتماع  |  الوظيفة الأولى: الرأي العام"),
              size=9.5, color="555555")
    hpPr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:color"), "AAB7C8")
    pbdr.append(bot)
    insert_ordered(hpPr, pbdr, PPR_AFTER_PBDR)

    fp = sec.footer.paragraphs[0]
    set_rtl(fp, WD_ALIGN_PARAGRAPH.CENTER)
    style_run(fp.add_run("صفحة "), size=9.5, color="555555")
    add_field(fp, "PAGE")
    style_run(fp.add_run(" من "), size=9.5, color="555555")
    add_field(fp, "NUMPAGES")
    style_run(fp.add_run("   —   جان شادي عبود ووديع هيثم سلوم"), size=9.5, color="555555")

    props = doc.core_properties
    props.title = "الوظيفة الأولى: الرأي العام — المدخل إلى علم الاجتماع"
    props.author = "جان شادي عبود (jan_436439) ووديع هيثم سلوم (wadee_440350)"
    props.subject = "وظيفة أكاديمية مشتركة — الجامعة الافتراضية السورية (SVU) — بكالوريوس تقانات المعلومات (BAIT)"
    props.keywords = "الرأي العام، علم الاجتماع، SVU، BAIT، السياسة الاقتصادية"

    # تحديث حقول أرقام الصفحات عند فتح الملف في Word
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    insert_ordered(settings, uf, SETTINGS_AFTER_UPDATE)

    jan_w, wadee_w = word_counts()
    jan_p, wadee_p = prose_counts()
    ar_refs, en_refs = sorted_refs()

    for b in BLOCKS:
        t = b["t"]
        if t == "cover_title":
            add_paragraph(doc, b["text"], size=b.get("size") or 16.5, bold=True,
                          color="1F4E79", align="center", indent=False,
                          space_after=5, line=1.1)
        elif t == "center":
            add_paragraph(doc, b["text"], size=b.get("size") or 11.5,
                          bold=b.get("bold", False), color=b.get("color"),
                          align="center", indent=False,
                          space_after=b.get("space_after", 6), line=1.15)
        elif t == "h1":
            add_paragraph(doc, b["text"], size=16, bold=True, color="1F4E79",
                          align="right", indent=False, space_before=10,
                          space_after=8, line=1.15, keep_next=True)
        elif t == "h2":
            add_paragraph(doc, b["text"], size=14, bold=True, color="2E74B5",
                          align="right", indent=False, space_before=10,
                          space_after=6, line=1.15, keep_next=True)
        elif t == "h3":
            add_paragraph(doc, b["text"], size=12.5, bold=True, color="1F3864",
                          align="right", indent=False, space_before=8,
                          space_after=4, line=1.15, keep_next=True)
        elif t == "p":
            add_paragraph(doc, b["text"], size=12.5,
                          align="justify", indent=b.get("indent", True))
        elif t in ("bullets", "numbers"):
            for i, item in enumerate(b["items"], 1):
                prefix = "•  " if t == "bullets" else ar_num(i) + ".  "
                para = doc.add_paragraph()
                set_rtl(para, WD_ALIGN_PARAGRAPH.JUSTIFY)
                pf = para.paragraph_format
                pf.left_indent = Cm(0.9)
                pf.first_line_indent = Cm(-0.55)
                pf.space_after = Pt(4)
                pf.line_spacing = 1.22
                style_run(para.add_run(prefix), size=12, bold=True, color="2E74B5")
                style_run(para.add_run(resolve(item)), size=12)
        elif t == "table":
            add_table(doc, b["headers"], b["rows"], b.get("widths"),
                      b.get("caption"))
        elif t == "toc":
            add_table(doc, ["م", "العنوان", "الطالب المسؤول", "الصفحة"],
                      b["rows"], [8, 54, 27, 11], fontsize=10)
        elif t == "note":
            add_box(doc, b["text"], b.get("label"), kind="note")
        elif t == "quote":
            add_box(doc, b["text"], b.get("source"), kind="quote")
        elif t == "hr":
            add_hr(doc)
        elif t == "spacer":
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(b.get("pts", 6))
            style_run(para.add_run(" "), size=2)
        elif t == "pb":
            doc.add_page_break()
        elif t == "signature":
            for line in b["lines"]:
                add_paragraph(doc, line, size=12, bold=True, color="1F4E79",
                              align="center", indent=False, space_after=4)
        elif t == "refs":
            add_paragraph(doc, "أولاً: المصادر والمراجع العربية", size=12.5,
                          bold=True, color="2E74B5", indent=False,
                          space_before=8, space_after=6, keep_next=True)
            for key, ref in ar_refs:
                para = doc.add_paragraph()
                set_rtl(para, WD_ALIGN_PARAGRAPH.RIGHT)
                pf = para.paragraph_format
                pf.left_indent = Cm(0.85)
                pf.first_line_indent = Cm(-0.85)
                pf.space_after = Pt(5)
                pf.line_spacing = 1.15
                style_run(para.add_run(ref["entry"]), size=10.5)
            add_paragraph(doc, "ثانياً: المراجع الأجنبية", size=12.5,
                          bold=True, color="2E74B5", indent=False,
                          space_before=10, space_after=6, keep_next=True)
            for key, ref in en_refs:
                para = doc.add_paragraph()
                set_rtl(para, WD_ALIGN_PARAGRAPH.RIGHT)
                pf = para.paragraph_format
                pf.left_indent = Cm(0.85)
                pf.first_line_indent = Cm(-0.85)
                pf.space_after = Pt(5)
                pf.line_spacing = 1.15
                style_run(para.add_run(ref["entry"]), size=10.5, font="Arial")
        elif t == "wordcount":
            rows = [
                ["جان شادي عبود", "jan_436439", "المقدمة + الأسئلة (1–3)",
                 f"{jan_p:,} ({jan_w:,} مع الجداول)", "≈ 5 صفحات"],
                ["وديع هيثم سلوم", "wadee_440350", "السؤالان (4–5) + الخاتمة",
                 f"{wadee_p:,} ({wadee_w:,} مع الجداول)", "≈ 5 صفحات"],
                ["المجموع", "—", "متن البحث كاملاً",
                 f"{jan_p + wadee_p:,} ({jan_w + wadee_w:,} مع الجداول)", "≈ 10 صفحات"],
            ]
            add_table(doc, ["الطالب", "الرقم الجامعي", "المهام المسندة",
                            "الكلمات التقريبية", "الحجم"],
                      rows, [21, 18, 29, 17, 15],
                      caption="جدول (6): توزيع المهام والجهد البحثي — إعداد الباحثين",
                      fontsize=10)
        else:
            raise ValueError("نوع كتلة غير معروف: " + t)

    doc.save(path)
    return path


# =====================================================================
# بناء Markdown
# =====================================================================

def md_escape(t):
    return t.replace("|", "\\|")


def md_table(headers, rows):
    out = ["| " + " | ".join(md_escape(str(h)) for h in headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        out.append("| " + " | ".join(md_escape(resolve(str(c))) for c in row) + " |")
    return "\n".join(out)


def build_markdown(path):
    jan_w, wadee_w = word_counts()
    jan_p, wadee_p = prose_counts()
    ar_refs, en_refs = sorted_refs()
    L = ["# الوظيفة الأولى: الرأي العام", "",
         "> مادة «المدخل إلى علم الاجتماع» — الجامعة الافتراضية السورية (SVU) — بكالوريوس تقانات المعلومات (BAIT).  ",
         "> إعداد: جان شادي عبود (`jan_436439`) ووديع هيثم سلوم (`wadee_440350`).  ",
         "> تاريخ التوزيع: 14/10/2026 — تاريخ الإعادة: 12/11/2026.", ""]
    for b in BLOCKS:
        t = b["t"]
        if t == "cover_title":
            L += ["## " + resolve(b["text"]), ""]
        elif t == "center":
            L += ["**" + resolve(b["text"]) + "**", ""]
        elif t == "h1":
            L += ["# " + resolve(b["text"]), ""]
        elif t == "h2":
            L += ["## " + resolve(b["text"]), ""]
        elif t == "h3":
            L += ["### " + resolve(b["text"]), ""]
        elif t == "p":
            L += [resolve(b["text"]), ""]
        elif t in ("bullets", "numbers"):
            for i, item in enumerate(b["items"], 1):
                L.append(("- " if t == "bullets" else f"{i}. ") + resolve(item))
            L.append("")
        elif t == "table":
            if b.get("caption"):
                L += ["**" + resolve(b["caption"]) + "**", ""]
            L += [md_table(b["headers"], b["rows"]), ""]
        elif t == "toc":
            L += [md_table(["م", "العنوان", "الطالب المسؤول", "الصفحة"], b["rows"]), ""]
        elif t == "note":
            L += ["> **" + b.get("label", "ملاحظة") + ":** " + resolve(b["text"]), ""]
        elif t == "quote":
            L += ["> " + resolve(b["text"]), ">", "> — *" + resolve(b.get("source") or "") + "*", ""]
        elif t == "hr":
            L += ["---", ""]
        elif t == "pb":
            L += ["<div style=\"page-break-after: always;\"></div>", ""]
        elif t == "spacer":
            L.append("")
        elif t == "signature":
            L += ["**" + resolve(x) + "**  " for x in b["lines"]] + [""]
        elif t == "refs":
            L += ["### أولاً: المصادر والمراجع العربية", ""]
            for key, ref in ar_refs:
                L += ["- " + ref["entry"], ""]
            L += ["### ثانياً: المراجع الأجنبية", ""]
            for key, ref in en_refs:
                L += ["- " + ref["entry"], ""]
        elif t == "wordcount":
            rows = [
                ["جان شادي عبود", "jan_436439", "المقدمة + الأسئلة (1–3)", f"{jan_p} ({jan_w} مع الجداول)", "≈ 5 صفحات"],
                ["وديع هيثم سلوم", "wadee_440350", "السؤالان (4–5) + الخاتمة", f"{wadee_p} ({wadee_w} مع الجداول)", "≈ 5 صفحات"],
                ["المجموع", "—", "متن البحث كاملاً", f"{jan_p + wadee_p} ({jan_w + wadee_w} مع الجداول)", "≈ 10 صفحات"],
            ]
            L += ["**جدول (6): توزيع المهام والجهد البحثي**", "",
                  md_table(["الطالب", "الرقم الجامعي", "المهام", "الكلمات", "الحجم"], rows), ""]
    io.open(path, "w", encoding="utf-8").write("\n".join(L))
    return path


# =====================================================================
# بناء HTML قابل للطباعة
# =====================================================================

CSS = """
:root{--ink:#1b2430;--blue:#1F4E79;--blue2:#2E74B5;--soft:#EAF1F8;--line:#cfdae8;}
*{box-sizing:border-box}
html{dir:rtl;lang:ar}
body{margin:0;padding:0;background:#eef1f5;color:var(--ink);
 font-family:'SVUAmiri','Traditional Arabic','Arabic Typesetting','Noto Naskh Arabic','Amiri',Arial,serif;
 font-size:13.2pt;line-height:1.85;}
.page{max-width:21cm;margin:1.2cm auto;padding:2.2cm 2.5cm;background:#fff;
 box-shadow:0 3px 18px rgba(0,0,0,.10);min-height:29.7cm;}
h1{font-size:19pt;color:var(--blue);margin:22px 0 12px;padding-bottom:7px;border-bottom:2px solid var(--blue);}
h2{font-size:16.5pt;color:var(--blue2);margin:20px 0 10px;}
h3{font-size:14.5pt;color:#1F3864;margin:16px 0 7px;}
p{margin:0 0 11px;text-align:justify;text-indent:.7cm;}
p.center{text-align:center;text-indent:0;}
p.small{font-size:11.5pt;}
.cover-title{text-align:center;color:var(--blue);font-weight:bold;font-size:18pt;margin:5px 0;}
hr{border:0;border-top:2px solid var(--blue);margin:14px 0;}
ul,ol{margin:0 0 12px;padding-inline-start:1.1cm;}
li{margin-bottom:7px;text-align:justify;}
table{width:100%;border-collapse:collapse;margin:12px 0 16px;font-size:11pt;}
th,td{border:1px solid var(--line);padding:6px 8px;text-align:right;vertical-align:top;}
th{background:#DCE6F1;color:var(--blue);font-weight:bold;text-align:center;}
tbody tr:nth-child(even){background:#F5F8FC;}
caption{caption-side:top;font-weight:bold;color:var(--blue);margin-bottom:6px;text-align:center;}
.box{background:var(--soft);border-right:5px solid var(--blue2);padding:12px 16px;margin:12px 0 16px;border-radius:3px;}
.box.quote{background:#F7F4EA;border-right-color:#BF9000;}
.box .lbl{font-weight:bold;color:var(--blue2);display:block;margin-bottom:4px;}
.box.quote .lbl{color:#8a6d00;}
.box p{margin:0;text-indent:0;}
.sig{text-align:center;font-weight:bold;color:var(--blue);}
.pb{break-before:page;page-break-before:always;}
.ref{font-size:11pt;padding-inline-start:.9cm;text-indent:-.9cm;margin-bottom:7px;text-align:right;}
.cite{color:#3b5b7a;}
@media print{
 body{background:#fff;font-size:12.6pt;}
 .page{margin:0;padding:0;box-shadow:none;max-width:none;min-height:0;}
 @page{size:A4;margin:2.2cm 2.4cm;}
 .pb{break-before:page;}
}
"""


def h(s):
    return html_lib.escape(resolve(str(s)))


def html_table(headers, rows, caption=None):
    out = ["<table>"]
    if caption:
        out.append("<caption>" + h(caption) + "</caption>")
    out.append("<thead><tr>" + "".join("<th>" + h(x) + "</th>" for x in headers) + "</tr></thead><tbody>")
    for row in rows:
        out.append("<tr>" + "".join("<td>" + h(x) + "</td>" for x in row) + "</tr>")
    out.append("</tbody></table>")
    return "\n".join(out)


def embedded_font_css():
    """تضمين خط عربي (Amiri) داخل ملف HTML لضمان ظهور التشكيل العربي في أي متصفح."""
    a = os.path.join(ROOT, "assets")
    out = []
    for fname, weight in (("amiri-regular.woff2", "400"), ("amiri-bold.woff2", "700")):
        f = os.path.join(a, fname)
        if not os.path.exists(f):
            return ""
        with open(f, "rb") as fh:
            data = base64.b64encode(fh.read()).decode("ascii")
        out.append("@font-face{font-family:'SVUAmiri';src:url(data:font/woff2;base64," +
                   data + ") format('woff2');font-weight:" + weight +
                   ";font-style:normal;font-display:swap;}")
    return "\n".join(out)


def build_html(path):
    jan_w, wadee_w = word_counts()
    jan_p, wadee_p = prose_counts()
    ar_refs, en_refs = sorted_refs()
    B = ['<!DOCTYPE html><html lang="ar" dir="rtl"><head><meta charset="utf-8">',
         '<meta name="viewport" content="width=device-width, initial-scale=1">',
         '<title>الوظيفة الأولى: الرأي العام — SVU</title>',
         "<style>" + embedded_font_css() + "\n" + CSS + "</style></head><body><div class=\"page\">"]
    for b in BLOCKS:
        t = b["t"]
        if t == "cover_title":
            s = b.get("size") or 18
            B.append(f'<div class="cover-title" style="font-size:{s}pt">{h(b["text"])}</div>')
        elif t == "center":
            B.append('<p class="center" style="font-size:' + str(b.get("size") or 12) + 'pt">' +
                     ("<b>" if b.get("bold") else "") + h(b["text"]) +
                     ("</b>" if b.get("bold") else "") + "</p>")
        elif t == "h1":
            B.append("<h1>" + h(b["text"]) + "</h1>")
        elif t == "h2":
            B.append("<h2>" + h(b["text"]) + "</h2>")
        elif t == "h3":
            B.append("<h3>" + h(b["text"]) + "</h3>")
        elif t == "p":
            B.append("<p>" + h(b["text"]) + "</p>")
        elif t in ("bullets", "numbers"):
            tag = "ul" if t == "bullets" else "ol"
            B.append("<" + tag + ">" + "".join("<li>" + h(x) + "</li>" for x in b["items"]) + "</" + tag + ">")
        elif t == "table":
            B.append(html_table(b["headers"], b["rows"], b.get("caption")))
        elif t == "toc":
            B.append(html_table(["م", "العنوان", "الطالب المسؤول", "الصفحة"], b["rows"]))
        elif t == "note":
            B.append('<div class="box"><span class="lbl">' + h(b.get("label", "")) +
                     '</span><p>' + h(b["text"]) + "</p></div>")
        elif t == "quote":
            B.append('<div class="box quote"><span class="lbl">' + h(b.get("source") or "") +
                     '</span><p>' + h(b["text"]) + "</p></div>")
        elif t == "hr":
            B.append("<hr>")
        elif t == "spacer":
            B.append('<div style="height:' + str(b.get("pts", 6)) + 'pt"></div>')
        elif t == "pb":
            B.append('<div class="pb"></div>')
        elif t == "signature":
            for line in b["lines"]:
                B.append('<p class="sig">' + h(line) + "</p>")
        elif t == "refs":
            B.append("<h3>أولاً: المصادر والمراجع العربية</h3>")
            for key, ref in ar_refs:
                B.append('<p class="ref">' + html_lib.escape(ref["entry"]) + "</p>")
            B.append("<h3>ثانياً: المراجع الأجنبية</h3>")
            for key, ref in en_refs:
                B.append('<p class="ref" dir="ltr" style="text-align:left">' +
                         html_lib.escape(ref["entry"]) + "</p>")
        elif t == "wordcount":
            rows = [
                ["جان شادي عبود", "jan_436439", "المقدمة + الأسئلة (1–3)", f"{jan_p} ({jan_w} مع الجداول)", "≈ 5 صفحات"],
                ["وديع هيثم سلوم", "wadee_440350", "السؤالان (4–5) + الخاتمة", f"{wadee_p} ({wadee_w} مع الجداول)", "≈ 5 صفحات"],
                ["المجموع", "—", "متن البحث كاملاً", f"{jan_p + wadee_p} ({jan_w + wadee_w} مع الجداول)", "≈ 10 صفحات"],
            ]
            B.append(html_table(["الطالب", "الرقم الجامعي", "المهام", "الكلمات", "الحجم"], rows,
                                "جدول (6): توزيع المهام والجهد البحثي"))
    B.append("</div></body></html>")
    io.open(path, "w", encoding="utf-8").write("\n".join(B))
    return path


# =====================================================================

def main():
    used, missing = validate()
    jan, wadee = word_counts()
    print("عدد المراجع:", len(REFERENCES), "| المستخدمة:", len(used), "| غير المُستشهد بها:", sorted(missing))
    jp, wp = prose_counts()
    print("كلمات القسم الأول (جان):", jp, "سردية -", jan, "شاملة")
    print("كلمات القسم الثاني (وديع):", wp, "سردية -", wadee, "شاملة")
    print("المجموع السردي:", jp + wp, "| الشامل:", jan + wadee)
    print("DOCX:", build_docx(OUT + ".docx"))
    print("MD  :", build_markdown(OUT + ".md"))
    print("HTML:", build_html(OUT + ".html"))


if __name__ == "__main__":
    main()
